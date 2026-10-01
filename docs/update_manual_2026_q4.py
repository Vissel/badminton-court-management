#!/usr/bin/env python3
"""
update_manual_2026_q4.py — Cập nhật tài liệu hướng dẫn sử dụng (.docx)
với các chức năng mới đã merge vào main từ nhánh develop-2026-Q4
(kèm thuê theo giờ + trả trước đang có trên main nhưng chưa được viết tài liệu).

Không bao gồm chức năng kho hàng (inventory) — chỉ có trên nhánh develop2027-Q1.

Cách chạy (từ thư mục docs/):
    python3 update_manual_2026_q4.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

DOC_PATH = Path(__file__).resolve().parent / \
    "Huong_dan_su_dung_He thong quan ly san cau long TC.docx"

HEADING_NUM_ID = "11"   # multilevel numbering used by Heading 1/2/3
BULLET_NUM_ID = "3"     # "-" bullet list used in the doc


# ── helpers ──────────────────────────────────────────────────────────────

def set_numpr(paragraph, num_id: str, ilvl: int = 0) -> None:
    pPr = paragraph._element.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl_el = OxmlElement("w:ilvl")
    ilvl_el.set(qn("w:val"), str(ilvl))
    numId_el = OxmlElement("w:numId")
    numId_el.set(qn("w:val"), num_id)
    numPr.append(ilvl_el)
    numPr.append(numId_el)
    pStyle = pPr.find(qn("w:pStyle"))
    if pStyle is not None:
        pStyle.addnext(numPr)
    else:
        pPr.insert(0, numPr)


def set_ind(paragraph, left: int) -> None:
    pPr = paragraph._element.get_or_add_pPr()
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    pPr.append(ind)


def set_spacing(paragraph, after: int = 0, line: int = 360) -> None:
    pPr = paragraph._element.get_or_add_pPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), str(after))
    spacing.set(qn("w:line"), str(line))
    spacing.set(qn("w:lineRule"), "auto")
    pPr.append(spacing)


def style_run(run) -> None:
    """Match existing body text: Arial 11pt, black."""
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_para(anchor, text: str, kind: str):
    """Insert a paragraph before `anchor`. kind: h2 | h3 | body | bullet."""
    p = anchor.insert_paragraph_before()
    if kind == "h2":
        p.style = "Heading 2"
        set_numpr(p, HEADING_NUM_ID, ilvl=1)
        p.add_run(text)
    elif kind == "h3":
        p.style = "Heading 3"
        set_numpr(p, HEADING_NUM_ID, ilvl=2)
        p.add_run(text)
    elif kind == "bullet":
        set_numpr(p, BULLET_NUM_ID, ilvl=0)
        set_spacing(p)
        style_run(p.add_run(text))
    else:  # body
        set_spacing(p)
        set_ind(p, 720)
        p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
        style_run(p.add_run(text))
    return p


# ── new content (end-user friendly, giọng miền Tây nhẹ) ──────────────────

NEW_SECTIONS = [
    ("h2", "Thuê sân theo giờ (khách thuê trọn sân)"),
    ("body",
     "Chức năng này dành cho khách thuê nguyên một sân theo giờ, hông cần chia đội "
     "như trận đấu thường. Tiền thuê được tính theo số giờ thuê (đơn giá lấy trong "
     "phần Cài đặt), cộng thêm tiền cầu nếu khách có xài cầu của sân."),
    ("body", "Cách bắt đầu cho thuê:"),
    ("body",
     "Bước 1: Rê chuột vô ô sân đang trống (sân chưa bắt đầu trận), bấm nút mũi tên "
     "nhỏ nằm kế nút “Bắt đầu” rồi chọn “Thuê theo giờ”."),
    ("body",
     "Bước 2: Trong hộp thoại “Thuê theo giờ”, gõ tên người thuê, chọn giờ bắt đầu "
     "với thời gian thuê (mấy giờ mấy phút)."),
    ("body",
     "Bước 3: Nếu khách mượn cầu thì bấm nút thêm cầu, chọn loại cầu với số lượng."),
    ("body", "Bước 4: Bấm “Bắt đầu” là xong nha."),
    ("body", "Trong lúc sân đang cho thuê:"),
    ("bullet",
     "Ô sân chạy đồng hồ đếm ngược thời gian còn lại, lúc này sân bị khóa nên hông "
     "xếp trận đấu được."),
    ("bullet", "Muốn đổi giờ thuê hay số cầu: bấm “Cập nhật” trên ô sân."),
    ("bullet",
     "Khách hông thuê nữa: bấm “Huỷ” rồi xác nhận, tiền thuê sẽ hông bị tính."),
    ("bullet",
     "Khách trả sân: bấm “Kết thúc” rồi xác nhận. Tiền thuê với tiền cầu sẽ vô bill "
     "của người thuê, mình thu tiền như bình thường."),

    ("h2", "Trả trước (khách ứng tiền lúc vô chơi)"),
    ("body",
     "Mỗi lần thêm người chơi mới vô phiên, hệ thống sẽ hỏi có thu tiền trước hay hông."),
    ("bullet",
     "Khách đưa tiền trước: gõ số tiền vô ô “Số tiền trả trước”, hoặc bấm nhanh mấy "
     "nút tiền có sẵn (10K, 20K…), rồi bấm “Xác nhận”."),
    ("bullet", "Hông thu trước: bấm “Bỏ qua” là xong."),
    ("body", "Lưu ý:"),
    ("bullet",
     "Tiền trả trước hiện trong danh sách dịch vụ của người chơi dưới tên “Trả trước” "
     "(ô màu xanh, ghi chữ “Đã trả”) — mục này hông xoá được nha."),
    ("bullet",
     "Lúc thanh toán, hệ thống tự trừ phần tiền đã trả trước, khách chỉ cần đưa phần "
     "còn lại thôi."),

    ("h2", "Quản lý nợ"),
    ("body",
     "Chức năng này giúp mình ghi nợ, theo dõi với thu nợ của người chơi."),

    ("h3", "Nợ được ghi như thế nào?"),
    ("bullet",
     "Tự động: khi phiên chơi bị đóng (cuối ngày, hoặc lần đăng nhập đầu tiên của "
     "ngày mới), người chơi nào chưa trả đủ tiền thì hệ thống tự ghi nợ phần còn "
     "thiếu, chi tiết dịch vụ nằm trong ghi chú của khoản nợ."),
    ("bullet",
     "Chủ động: mình ghi nợ ngay lúc thanh toán cho khách (xem mục bên dưới)."),

    ("h3", "Ghi nợ với thu nợ ngay lúc thanh toán"),
    ("body", "Trong hộp thoại thanh toán của người chơi:"),
    ("bullet",
     "Muốn ghi nợ: bấm “+ Ghi nợ”, gõ số tiền khách nợ lại với ghi chú nếu cần. Khách "
     "chỉ đưa phần bill trừ ra số tiền nợ, phần nợ được lưu lại để thu sau."),
    ("bullet",
     "Muốn thu luôn nợ cũ: nếu khách đang có nợ, trong hộp thoại sẽ thấy dòng "
     "“Nợ: số tiền”. Bấm vô đó để xem từng khoản, tick chọn khoản muốn thu rồi bấm "
     "“Thêm vào thanh toán” — tiền nợ cũ sẽ cộng luôn vô bill hiện tại, khách trả "
     "một lần cho gọn."),
    ("bullet", "Chọn hình thức trả tiền: “Tiền mặt” hay “Chuyển khoản”."),

    ("h3", "Trang Quản lý nợ"),
    ("body", "Bấm “Quản lý nợ” trên thanh menu phía trên của màn hình."),
    ("body", "Tab “Nợ hiện tại” — danh sách người đang còn nợ:"),
    ("bullet", "Mỗi dòng là một người nợ, kèm số khoản nợ."),
    ("bullet",
     "Bấm vô dòng tên để xổ xuống, xem từng khoản nợ: ngày nợ, số tiền, ghi chú."),
    ("bullet",
     "Để thu nợ: tick chọn các khoản muốn thu rồi bấm “Thu nợ”. Cũng có thể bấm "
     "“Thu nợ” luôn, trong hộp thoại gõ số tiền muốn thu — hệ thống tự chia ra từng "
     "khoản nợ. Xong chọn “Tiền mặt”/“Chuyển khoản”, ghi chú nếu cần, rồi bấm "
     "“Thanh toán”."),
    ("body", "Tab “Lịch sử” — xem lại các khoản nợ đã trả hay chưa:"),
    ("bullet",
     "Cột “Đã trả / Tổng” cho biết người đó đã trả được bao nhiêu phần nợ."),
    ("bullet",
     "Muốn xem theo thời gian: chọn “Từ ngày”, “Đến ngày” rồi bấm “Lọc”; muốn xem "
     "hết thì bấm “Tất cả”."),
    ("bullet",
     "Bấm vô dòng người chơi để xem chi tiết từng khoản với trạng thái “Đã trả”, "
     "“Trả một phần” hay “Chưa trả”. Trong màn hình chi tiết có nút “Excel” để xuất "
     "lịch sử nợ riêng của người đó."),
    ("body", "Mấy tiện ích khác:"),
    ("bullet",
     "Ô “Tìm theo tên người chơi”: gõ tên để tìm nhanh, gõ không dấu cũng được luôn "
     "(gõ “nguyen” vẫn ra “Nguyễn”)."),
    ("bullet",
     "Nút “Xuất Excel”: tải về máy file Excel danh sách nợ đang hiển thị trên màn hình."),
    ("bullet",
     "Dưới bảng có phân trang, mình chọn xem 10, 20 hay 50 dòng một trang."),
]

# rows appended to the function-scope table (STT | Nhóm chức năng | Nội dung | Nhóm người dùng)
NEW_TABLE_ROWS = [
    ("10", "Thuê sân theo giờ", "Cho thuê trọn sân, tính tiền theo thời gian thuê",
     "Quản trị viên / Thu ngân"),
    ("11", "Trả trước", "Ghi nhận tiền khách ứng trước khi vào phiên",
     "Quản trị viên / Thu ngân"),
    ("12", "Quản lý nợ", "Ghi nợ, theo dõi, thu nợ và xuất báo cáo nợ",
     "Quản trị viên / Thu ngân, Chủ sân"),
]

# bullet added into the "Khi đóng phiên, hệ thống sẽ:" list
AUTO_DEBIT_BULLET = (
    "Người chơi nào chưa trả đủ tiền sẽ được hệ thống tự động ghi nợ phần còn "
    "thiếu (xem thêm ở phần Quản lý nợ)."
)


def main() -> None:
    if not DOC_PATH.exists():
        print(f"Không thấy file: {DOC_PATH}")
        sys.exit(1)

    doc = Document(DOC_PATH)

    # 1) Append rows to the function-scope table
    table = doc.tables[0]
    for stt, group, work, owners in NEW_TABLE_ROWS:
        cells = table.add_row().cells
        for cell, text in zip(cells, (stt, group, work, owners)):
            cell.text = ""
            style_run(cell.paragraphs[0].add_run(text))
    print(f"Đã thêm {len(NEW_TABLE_ROWS)} dòng vào bảng phạm vi chức năng.")

    # 2) Add auto-debit bullet into the session-close list
    anchor_session = None
    for p in doc.paragraphs:
        if p.text.strip() == "Quản lý người trong phiên" and p.style.name == "Heading 2":
            anchor_session = p
            break
    if anchor_session is None:
        print("Cảnh báo: không tìm thấy mục 'Quản lý người trong phiên'.")
    else:
        bullet = anchor_session.insert_paragraph_before()
        set_numpr(bullet, BULLET_NUM_ID, ilvl=0)
        set_spacing(bullet, line=276)
        bullet.add_run(AUTO_DEBIT_BULLET)
        print("Đã thêm ghi chú tự động ghi nợ vào mục đóng phiên.")

    # 3) Insert new feature sections right before the "Lưu ý" Heading 1
    anchor_note = None
    for p in doc.paragraphs:
        if p.text.strip() == "Lưu ý" and p.style.name == "Heading 1":
            anchor_note = p
            break
    if anchor_note is None:
        print("Cảnh báo: không tìm thấy mục 'Lưu ý' — sẽ thêm nội dung ở cuối tài liệu.")
        # fallback: append at end using a dummy anchor paragraph, then move it away
        anchor_note = doc.add_paragraph()
        for kind, text in NEW_SECTIONS:
            add_para(anchor_note, text, kind)
        anchor_note._element.getparent().remove(anchor_note._element)
    else:
        for kind, text in NEW_SECTIONS:
            add_para(anchor_note, text, kind)
        print(f"Đã chèn {len(NEW_SECTIONS)} đoạn nội dung mới trước mục 'Lưu ý'.")

    doc.save(DOC_PATH)
    print(f"Xong! Đã cập nhật: {DOC_PATH}")


if __name__ == "__main__":
    main()
