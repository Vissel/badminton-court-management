#!/usr/bin/env python3
"""
billing_guide_to_docx.py — Convert a billing user-guide Markdown file to .docx.

Extends the conventions in generate_manual_docx.py with support for:
pipe tables, **bold** inline, `code` inline, and > blockquotes.

Usage (from docs/):
    python3 billing_guide_to_docx.py [--source <md>] [--output <docx>]
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt

BOLD_PATTERN = re.compile(r"\*\*(.+?)\*\*")
CODE_PATTERN = re.compile(r"`([^`]+)`")
TABLE_ROW = re.compile(r"^\|.*\|$")
TABLE_SEP = re.compile(r"^\|[\s:|-]+\|$")


def set_font(style, name: str = "Arial", size=None) -> None:
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        style.font.size = size


def add_inline(paragraph, text: str) -> None:
    """Add runs to a paragraph, honouring **bold** and `code` spans."""
    pos = 0
    for match in re.finditer(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(BOLD_PATTERN.fullmatch(token).group(1))
            run.bold = True
        else:
            run = paragraph.add_run(CODE_PATTERN.fullmatch(token).group(1))
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(10)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def add_table(doc, rows: list[str]) -> None:
    cells = [split_table_row(r) for r in rows]
    width = max(len(r) for r in cells)
    table = doc.add_table(rows=len(cells), cols=width)
    table.style = "Table Grid"
    for i, row in enumerate(cells):
        for j, cell_text in enumerate(row):
            cell = table.cell(i, j)
            cell.paragraphs[0].clear() if False else None  # keep default para
            add_inline(cell.paragraphs[0], cell_text)
            if i == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    doc.add_paragraph("")


def build_docx(md_path: Path, docx_path: Path) -> Path:
    if not md_path.exists():
        print(f"Error: source markdown not found: {md_path}")
        sys.exit(1)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()

    for name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3"]:
        set_font(doc.styles[name])
    doc.styles["Normal"].font.size = Pt(11)

    number_buffer: list[str] = []
    table_buffer: list[str] = []

    def flush_number_buffer() -> None:
        nonlocal number_buffer
        for item in number_buffer:
            add_inline(doc.add_paragraph(style="List Number"), item)
        number_buffer = []

    def flush_table_buffer() -> None:
        nonlocal table_buffer
        rows = [r for r in table_buffer if not TABLE_SEP.match(r)]
        if rows:
            add_table(doc, rows)
        table_buffer = []

    for raw in lines:
        line = raw.rstrip()
        stripped = line.strip()

        if TABLE_ROW.match(stripped):
            flush_number_buffer()
            table_buffer.append(stripped)
            continue
        flush_table_buffer()

        if not stripped:
            flush_number_buffer()
            continue

        if line.startswith("#### "):
            flush_number_buffer()
            doc.add_paragraph(line[5:].strip(), style="Heading 3")
        elif line.startswith("### "):
            flush_number_buffer()
            doc.add_paragraph(line[4:].strip(), style="Heading 2")
        elif line.startswith("## "):
            flush_number_buffer()
            doc.add_paragraph(line[3:].strip(), style="Heading 1")
        elif line.startswith("# "):
            flush_number_buffer()
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            p.add_run(line[2:].strip())
        elif stripped.startswith("> "):
            flush_number_buffer()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            for run_holder in [p]:
                run = run_holder.add_run(stripped[2:].strip())
                run.italic = True
        elif line.lstrip().startswith("- "):
            flush_number_buffer()
            add_inline(doc.add_paragraph(style="List Bullet"),
                       line.lstrip()[2:].strip())
        else:
            parts = stripped.split(". ", 1)
            if len(parts) == 2 and parts[0].isdigit():
                number_buffer.append(parts[1])
                continue
            flush_number_buffer()
            para = doc.add_paragraph()
            para.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
            add_inline(para, stripped)

    flush_number_buffer()
    flush_table_buffer()

    for p in doc.paragraphs:
        if p.style.name != "Title":
            p.paragraph_format.space_after = Pt(6)

    doc.save(docx_path)
    return docx_path


def main() -> None:
    docs_dir = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser()
    parser.add_argument("--source",
                        default=str(docs_dir / "billing-user-guide-vi.md"))
    parser.add_argument("--output",
                        default=str(docs_dir / "billing-user-guide-vi.docx"))
    args = parser.parse_args()
    result = build_docx(Path(args.source), Path(args.output))
    print(f"Done: {result}")


if __name__ == "__main__":
    main()
